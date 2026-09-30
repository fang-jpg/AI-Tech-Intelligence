from openpyxl import load_workbook

from ai_tech_intelligence.reports import export_reports


def test_report_exports_expected_workbook(tmp_path) -> None:
    rows = [
        {
            "company_name": "OpenAI",
            "people_json": [{"name": "Sam Altman", "role": "CEO"}],
            "source_name": "OpenAI Newsroom",
            "published_at": "2026-09-01",
            "source_type": "official_blog",
            "title": "Test title",
            "url": "https://openai.com/index/test",
            "core_points_json": ["Point one"],
            "technical_signals_json": [{"tag": "agent", "evidence": "agent"}],
            "attributed_statements_json": [{"speaker": "Sam Altman", "statement": "Test", "evidence": "Test"}],
            "trend_summary": "基于公开信号的推断：Agent 能力可能继续增强。",
            "evidence_level": "Strong inference",
            "confidence": 0.72,
            "time_window": "未来 6-12 个月",
            "tier": "A",
            "official": 1,
            "attribution_status": "named_person_mention",
            "grounding_passed": 1,
            "analysis_method": "llm_two_stage",
            "model": "test-model",
            "fetched_at": "2026-09-18T00:00:00+00:00",
            "analyzed_at": "2026-09-18T00:01:00+00:00",
        }
    ]
    outputs = export_reports(rows, tmp_path)
    workbook = load_workbook(outputs["xlsx"], data_only=False)
    assert workbook.sheetnames == ["情报明细", "厂商趋势", "说明与口径"]
    assert workbook["情报明细"]["A2"].value == "OpenAI"
    assert workbook["情报明细"]["G2"].hyperlink.target == "https://openai.com/index/test"
    assert outputs["csv"].exists()
    assert outputs["markdown"].exists()
