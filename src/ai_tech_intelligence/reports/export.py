from __future__ import annotations

import csv
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from ..taxonomy import (
    adoption_stage,
    classify_record,
    infer_maturity,
    load_technology_taxonomy,
    matched_modules,
)


DETAIL_HEADERS = [
    "模型厂商",
    "高管/来源",
    "职位",
    "日期",
    "来源类型",
    "原文标题",
    "高管发言/博客链接",
    "核心观点",
    "技术标签",
    "AI未来模型动向分析",
    "证据等级",
    "置信度",
    "时间窗口",
    "来源等级",
    "是否官方",
    "归因状态",
    "证据落地校验",
    "分析方法/模型",
    "抓取时间",
]


def _join_points(values: list[Any]) -> str:
    return "\n".join(f"{index}. {value}" for index, value in enumerate(values, 1))


def _signal_tags(values: list[dict[str, Any]]) -> str:
    tags = [str(item.get("tag", "")).strip() for item in values if str(item.get("tag", "")).strip()]
    return ", ".join(dict.fromkeys(tags))


def _detail_row(item: dict[str, Any]) -> list[Any]:
    people = item.get("people_json") or []
    statements = item.get("attributed_statements_json") or []
    speaker_names = list(dict.fromkeys(str(statement.get("speaker", "")).strip() for statement in statements if str(statement.get("speaker", "")).strip()))
    role_by_name = {str(person.get("name", "")): str(person.get("role", "")) for person in people}
    default_source = item.get("source_name", "")
    if item.get("official") and item.get("source_type") == "search_discovery":
        default_source = f"{item.get('company_name', '')} 官方来源"
    people_name = "、".join(speaker_names) or default_source
    people_role = "、".join(role_by_name.get(name, "待核验") for name in speaker_names)
    published = item.get("published_at")
    try:
        published_value: date | str | None = date.fromisoformat(published) if published else None
    except ValueError:
        published_value = published
    confidence = item.get("confidence")
    confidence_value = float(confidence) if confidence is not None else None
    model = " / ".join(value for value in [item.get("analysis_method"), item.get("model")] if value)
    return [
        item.get("company_name", ""),
        people_name,
        people_role,
        published_value,
        item.get("source_type", ""),
        item.get("title", ""),
        item.get("url", ""),
        _join_points(item.get("core_points_json") or []),
        _signal_tags(item.get("technical_signals_json") or []),
        item.get("trend_summary") or "待分析",
        item.get("evidence_level") or "待分析",
        confidence_value,
        item.get("time_window") or "",
        item.get("tier", ""),
        "是" if item.get("official") else "否",
        item.get("attribution_status", ""),
        "通过" if item.get("grounding_passed") else ("未通过" if item.get("analyzed_at") else "待分析"),
        model or "待分析",
        item.get("fetched_at", ""),
    ]


def export_reports(rows: list[dict[str, Any]], output_dir: Path, prefix: str = "ai_executive_intelligence") -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = output_dir / f"{prefix}_{stamp}"
    detail_rows = [_detail_row(item) for item in rows]
    xlsx_path = base.with_suffix(".xlsx")
    csv_path = base.with_suffix(".csv")
    md_path = base.with_suffix(".md")
    _write_xlsx(xlsx_path, rows, detail_rows)
    _write_csv(csv_path, detail_rows)
    _write_markdown(md_path, detail_rows)
    return {"xlsx": xlsx_path, "csv": csv_path, "markdown": md_path}


def _write_xlsx(path: Path, source_rows: list[dict[str, Any]], detail_rows: list[list[Any]]) -> None:
    workbook = Workbook()
    detail = workbook.active
    detail.title = "情报明细"
    summary = workbook.create_sheet("厂商趋势")
    classification = workbook.create_sheet("前沿技术分类")
    evolution = workbook.create_sheet("技术演进沙盘")
    profiles = workbook.create_sheet("主体战略画像")
    notes = workbook.create_sheet("说明与口径")

    detail.append(DETAIL_HEADERS)
    for row in detail_rows:
        detail.append(row)
    _style_detail_sheet(detail)

    summary_headers = ["模型厂商", "来源数", "已分析数", "最新日期", "主要技术标签", "最新趋势判断", "平均置信度"]
    summary.append(summary_headers)
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in source_rows:
        grouped[str(item.get("company_name", ""))].append(item)
    for company in sorted(grouped):
        items = grouped[company]
        analyzed = [item for item in items if item.get("analyzed_at")]
        dates = [str(item.get("published_at")) for item in items if item.get("published_at")]
        summary_tags = _normalized_summary_tags(analyzed)
        latest_analyzed = _representative_model_signal(analyzed)
        confidences = [float(item["confidence"]) for item in analyzed if item.get("confidence") is not None]
        summary.append(
            [
                company,
                len(items),
                len(analyzed),
                max(dates) if dates else "",
                ", ".join(summary_tags),
                latest_analyzed.get("trend_summary", "待分析"),
                sum(confidences) / len(confidences) if confidences else None,
            ]
        )
    _style_summary_sheet(summary)

    classified = _technology_classification_rows(source_rows)
    _write_classification_sheet(classification, classified)
    _write_evolution_sheet(evolution, classified)
    _write_profile_sheet(profiles, source_rows, classified)

    notes_rows = [
        ["字段", "说明"],
        ["范围", "国内外关键模型厂商的第一方博客、研究页面、官方活动、正式演讲及可核验采访。"],
        ["Explicit", "原文明确给出未来计划、承诺或下一步方向。"],
        ["Strong inference", "多个一致信号共同指向某一方向，但原文未直接承诺。"],
        ["Speculative", "仅有单一或模糊信号，必须按低置信度理解。"],
        ["官方团队内容", "官方页面没有明确出现已登记高管姓名，不能视为个人发言。"],
        ["证据落地校验", "分析所引用证据短句能否在抓取原文中找到。失败项应人工复核。"],
        ["限制", "AI未来模型动向为基于公开信号的推断，不是厂商已确认事实；网页抓取也可能受动态渲染、登录或反爬限制。"],
    ]
    for row in notes_rows:
        notes.append(row)
    _style_notes_sheet(notes)

    workbook.save(path)


def _style_header(sheet: Any, width: int) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(name="Arial", color="FFFFFF", bold=True, size=10)
    for cell in sheet[1][:width]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.sheet_view.showGridLines = False


def _style_detail_sheet(sheet: Any) -> None:
    _style_header(sheet, len(DETAIL_HEADERS))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [16, 22, 26, 12, 18, 42, 48, 50, 26, 60, 18, 12, 18, 12, 12, 22, 16, 24, 24]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[3].number_format = "yyyy-mm-dd"
        row[11].number_format = "0.0%"
        if row[6].value:
            row[6].hyperlink = str(row[6].value)
            row[6].style = "Hyperlink"
        sheet.row_dimensions[row[0].row].height = 96
    if sheet.max_row >= 2:
        sheet.conditional_formatting.add(
            f"Q2:Q{sheet.max_row}",
            FormulaRule(formula=['$Q2="未通过"'], fill=PatternFill("solid", fgColor="FCE8E6")),
        )
        sheet.conditional_formatting.add(
            f"Q2:Q{sheet.max_row}",
            FormulaRule(formula=['$Q2="通过"'], fill=PatternFill("solid", fgColor="E6F4EA")),
        )
    if sheet.max_row >= 2:
        table = Table(displayName="IntelDetailTable", ref=f"A1:S{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)


def _style_summary_sheet(sheet: Any) -> None:
    _style_header(sheet, 7)
    sheet.freeze_panes = "A2"
    widths = [18, 12, 12, 14, 38, 72, 14]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[6].number_format = "0.0%"
        sheet.row_dimensions[row[0].row].height = 60
    if sheet.max_row >= 2:
        table = Table(displayName="CompanyTrendTable", ref=f"A1:G{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def _style_notes_sheet(sheet: Any) -> None:
    _style_header(sheet, 2)
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 100
    light_border = Border(bottom=Side(style="thin", color="D9E2F3"))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = light_border
        sheet.row_dimensions[row[0].row].height = 36


CLASSIFICATION_HEADERS = [
    "信息源类型", "主体", "发言人/来源", "职位", "日期", "一级技术分类", "技术模块/命中词",
    "是否主分类", "证据摘要", "趋势/主张", "成熟度 1-5", "成熟度说明", "采用阶段", "证据等级",
    "置信度", "时间窗口", "官方链接", "内容基础", "归因/证据状态", "分析模型",
]


def _technology_classification_rows(source_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for item in source_rows:
        statements = item.get("attributed_statements_json") or []
        people = item.get("people_json") or []
        speaker_names = list(dict.fromkeys(str(value.get("speaker", "")).strip() for value in statements if str(value.get("speaker", "")).strip()))
        role_by_name = {str(person.get("name", "")): str(person.get("role", "")) for person in people}
        source_name = "、".join(speaker_names) or str(item.get("source_name", ""))
        source_kind = "企业高管" if speaker_names else "厂商官方来源"
        role = "、".join(role_by_name.get(name, "待核验") for name in speaker_names)
        categories = classify_record(item, max_categories=3)
        if not categories:
            categories = [{"id": "other", "name": "其他前沿技术与应用", "module_terms": ["应用与交叉方向"], "matched_keywords": []}]
        maturity, maturity_label = infer_maturity(item)
        evidence = next((str(value.get("evidence", "")).strip() for value in statements if str(value.get("evidence", "")).strip()), "")
        if not evidence:
            evidence = next((str(value).strip() for value in item.get("core_points_json") or [] if str(value).strip()), str(item.get("title", "")))
        text = " ".join([str(item.get("title", "")), _signal_tags(item.get("technical_signals_json") or []), evidence, str(item.get("trend_summary", ""))])
        for index, category in enumerate(categories):
            output.append(
                {
                    "source_kind": source_kind,
                    "subject": str(item.get("company_name", "")),
                    "source_name": source_name,
                    "role": role,
                    "date": item.get("published_at"),
                    "category_id": category["id"],
                    "category": category["name"],
                    "module": "、".join(category.get("matched_keywords", [])[:4]) or "、".join(matched_modules(category, text)),
                    "primary": index == 0,
                    "evidence": evidence,
                    "trend": str(item.get("trend_summary", "")),
                    "maturity": maturity,
                    "maturity_label": maturity_label,
                    "adoption": adoption_stage(item, maturity),
                    "evidence_level": str(item.get("evidence_level", "")),
                    "confidence": item.get("confidence"),
                    "time_window": str(item.get("time_window", "")),
                    "url": str(item.get("url", "")),
                    "content_basis": str(item.get("source_type", "")),
                    "evidence_status": str(item.get("attribution_status", "")) + (" / 证据通过" if item.get("grounding_passed") else " / 待复核"),
                    "model": " / ".join(value for value in [item.get("analysis_method"), item.get("model")] if value),
                    "product_focus": category.get("product_focus", ""),
                }
            )
    return output


def _write_classification_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    sheet.append(CLASSIFICATION_HEADERS)
    for item in rows:
        published = item["date"]
        try:
            published = date.fromisoformat(published) if published else None
        except ValueError:
            pass
        sheet.append([
            item["source_kind"], item["subject"], item["source_name"], item["role"], published,
            item["category"], item["module"], "是" if item["primary"] else "否", item["evidence"], item["trend"],
            item["maturity"], item["maturity_label"], item["adoption"], item["evidence_level"], item["confidence"],
            item["time_window"], item["url"], item["content_basis"], item["evidence_status"], item["model"],
        ])
    widths = [16, 18, 22, 22, 12, 24, 28, 12, 52, 60, 12, 20, 20, 18, 12, 18, 48, 18, 24, 28]
    _style_new_analysis_sheet(sheet, widths, "TechnologyClassificationTable", 82)
    for row in sheet.iter_rows(min_row=2):
        row[4].number_format = "yyyy-mm-dd"
        row[14].number_format = "0.0%"
        if row[16].value:
            row[16].hyperlink = str(row[16].value)
            row[16].style = "Hyperlink"


def _write_evolution_sheet(sheet: Any, rows: list[dict[str, Any]]) -> None:
    headers = ["前沿技术", "核心模块/演进方向", "技术状态", "采用阶段", "信号数", "主体数", "代表主体", "代表证据", "未来可观察信号", "产品/战略影响", "证据等级", "平均置信度", "主要来源"]
    sheet.append(headers)
    taxonomy = {item["id"]: item for item in load_technology_taxonomy()}
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in rows:
        grouped[item["category_id"]].append(item)
    for category_id, items in sorted(grouped.items(), key=lambda pair: next((index for index, value in enumerate(load_technology_taxonomy()) if value["id"] == pair[0]), 999)):
        category = taxonomy.get(category_id, {"name": "其他前沿技术与应用", "product_focus": "建议人工复核并补充分类规则。"})
        ranked = sorted(items, key=lambda value: (float(value.get("confidence") or 0.0), int(value["maturity"]), value["primary"]), reverse=True)
        confidences = [float(item["confidence"]) for item in items if item.get("confidence") is not None]
        subjects = list(dict.fromkeys(item["subject"] for item in ranked))
        modules = [name for name, _ in Counter(part for item in items for part in str(item["module"]).split("、") if part).most_common(5)]
        adoption = " / ".join(name for name, _ in Counter(item["adoption"] for item in items).most_common(2))
        avg_maturity = round(sum(int(item["maturity"]) for item in items) / len(items))
        sheet.append([
            category["name"], "、".join(modules), f"{avg_maturity}级：{ranked[0]['maturity_label']}（最高信号）", adoption,
            len(items), len(set(subjects)), "、".join(subjects[:8]),
            "\n".join(f"{item['subject']}：{item['evidence'][:180]}" for item in ranked[:3]),
            "\n".join([item["trend"] for item in ranked if item["trend"]][:2]) or "关注正式模型、数据集、评测或产品发布。",
            category.get("product_focus", ""),
            " / ".join(name for name, _ in Counter(item["evidence_level"] for item in items if item["evidence_level"]).most_common(2)),
            sum(confidences) / len(confidences) if confidences else None,
            "\n".join(dict.fromkeys(item["url"] for item in ranked[:4] if item["url"])),
        ])
    _style_new_analysis_sheet(sheet, [24, 34, 24, 24, 10, 10, 48, 72, 72, 64, 20, 14, 56], "TechnologyEvolutionTable", 108)
    for row in sheet.iter_rows(min_row=2):
        row[11].number_format = "0.0%"
        if row[12].value:
            row[12].hyperlink = str(row[12].value).splitlines()[0]
            row[12].style = "Hyperlink"


def _write_profile_sheet(sheet: Any, source_rows: list[dict[str, Any]], classified: list[dict[str, Any]]) -> None:
    headers = ["主体", "类型", "核心定位", "研究路线", "产品/应用信号", "组织/资本/基础设施", "关键行动", "未来重点", "领先指标", "风险/反向信号", "证据等级", "平均置信度", "主要来源"]
    sheet.append(headers)
    grouped_source: dict[str, list[dict[str, Any]]] = defaultdict(list)
    grouped_classified: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for item in source_rows:
        grouped_source[str(item.get("company_name", ""))].append(item)
    for item in classified:
        if item["primary"]:
            grouped_classified[item["subject"]].append(item)
    for subject in sorted(grouped_source):
        source_items = grouped_source[subject]
        class_items = grouped_classified[subject]
        categories = [name for name, _ in Counter(item["category"] for item in class_items).most_common(4)]
        signals_text = " ".join(str(item.get("core_points_json") or "") for item in source_items)
        product_points = [str(point) for item in source_items for point in item.get("core_points_json") or [] if re.search(r"产品|应用|平台|服务|用户|客户|API|App|部署|上线", str(point), re.I)]
        infra_points = [str(point) for item in source_items for point in item.get("core_points_json") or [] if re.search(r"资本|投资|组织|部门|芯片|算力|服务器|数据中心|基础设施|cloud|GPU|XPU", str(point), re.I)]
        representative = _representative_model_signal(source_items)
        directions = representative.get("trend_directions_json") or []
        indicators = [str((value.get("quantitative_outlook") or {}).get("observable_metric", "")) for value in directions if isinstance(value, dict)]
        falsifiers = [str(item) for value in directions if isinstance(value, dict) for item in (value.get("quantitative_outlook") or {}).get("falsifiers", [])]
        confidences = [float(item["confidence"]) for item in source_items if item.get("confidence") is not None]
        full_stack = {"LLM架构与语言能力", "Agent与工具使用", "训练推理效率与基础设施"}.issubset(set(categories))
        positioning = "基础模型、Agent与基础设施协同推进" if full_stack else f"重点围绕{'、'.join(categories[:3])}形成技术组合"
        sheet.append([
            subject, "企业/高管", positioning, "\n".join(categories),
            "\n".join(dict.fromkeys(product_points[:3])) or "现有材料未形成明确的产品/应用信号。",
            "\n".join(dict.fromkeys(infra_points[:3])) or "现有材料未形成可核验的组织、资本或基础设施动作信号。",
            "\n".join(f"{item.get('published_at') or '日期待核验'}：{item.get('title', '')}" for item in sorted(source_items, key=lambda value: str(value.get("published_at") or ""), reverse=True)[:3]),
            str(representative.get("trend_summary", "现有材料不足以形成主体级趋势判断。")),
            "\n".join(value for value in indicators if value) or f"关注{'、'.join(categories[:3])}方向的正式发布、评测或产品落地。",
            "\n".join(value for value in falsifiers if value) or "若后续正式发布、产品采用或评测结果未延续当前信号，应下调判断。",
            str(representative.get("evidence_level", "")), sum(confidences) / len(confidences) if confidences else None,
            "\n".join(dict.fromkeys(str(item.get("url", "")) for item in source_items[:4] if item.get("url"))),
        ])
    _style_new_analysis_sheet(sheet, [20, 14, 34, 34, 64, 64, 58, 72, 54, 54, 18, 14, 56], "EntityStrategyProfileTable", 108)
    for row in sheet.iter_rows(min_row=2):
        row[11].number_format = "0.0%"
        if row[12].value:
            row[12].hyperlink = str(row[12].value).splitlines()[0]
            row[12].style = "Hyperlink"


def _style_new_analysis_sheet(sheet: Any, widths: list[int], table_name: str, row_height: int) -> None:
    _style_header(sheet, len(widths))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="1F2937")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=Side(style="hair", color="D9E2F3"))
        sheet.row_dimensions[row[0].row].height = row_height
    if sheet.max_row >= 2:
        table = Table(displayName=table_name, ref=f"A1:{get_column_letter(len(widths))}{sheet.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def _write_csv(path: Path, rows: list[list[Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(DETAIL_HEADERS)
        for row in rows:
            writer.writerow([value.isoformat() if isinstance(value, date) else value for value in row])


def _write_markdown(path: Path, rows: list[list[Any]]) -> None:
    selected = [0, 1, 3, 6, 7, 9, 10, 11]
    headers = [DETAIL_HEADERS[index] for index in selected]
    lines = ["# AI 模型厂商高管技术情报", "", "| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"]
    for row in rows:
        values = []
        for index in selected:
            value = row[index]
            if isinstance(value, date):
                value = value.isoformat()
            if isinstance(value, float):
                value = f"{value:.0%}"
            value = str(value or "").replace("\n", "<br>").replace("|", "\\|")
            values.append(value)
        lines.append("| " + " | ".join(values) + " |")
    lines.extend(["", "> AI 未来模型动向是基于公开信号的推断，不是厂商已确认事实。"])
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


SUMMARY_TAG_RULES: list[tuple[str, tuple[str, ...]]] = [
    ("Agent", ("agent", "智能体", "tool use", "工具调用", "computer use")),
    ("Coding", ("coding", "code", "编程", "代码", "软件工程")),
    ("Reasoning", ("reasoning", "推理", "强化学习", "reinforcement")),
    ("Multimodal", ("multimodal", "多模态", "视觉", "vision", "audio", "video")),
    ("Long Context", ("context", "上下文", "长文本")),
    ("Efficiency / Infra", ("efficiency", "inference", "training", "moe", "sparse", "效率", "推理", "训练", "芯片", "服务器", "基础设施")),
    ("Open Source", ("open source", "open-source", "开源", "权重")),
    ("AI for Science", ("science", "scientific", "科学")),
    ("Safety", ("safety", "安全", "governance", "治理")),
    ("World / Embodied", ("world model", "世界模型", "robot", "机器人", "具身")),
]


def _normalized_summary_tags(items: list[dict[str, Any]]) -> list[str]:
    counts: Counter[str] = Counter()
    raw_tags: Counter[str] = Counter()
    for item in items:
        for signal in item.get("technical_signals_json") or []:
            raw = " ".join(str(signal.get(key, "")) for key in ("tag", "description")).casefold()
            tag = str(signal.get("tag", "")).strip()
            if tag:
                raw_tags[tag] += 1
            for label, keywords in SUMMARY_TAG_RULES:
                if any(keyword.casefold() in raw for keyword in keywords):
                    counts[label] += 1
    if counts:
        return [label for label, _ in counts.most_common(6)]
    return [label for label, _ in raw_tags.most_common(6)]


def _representative_model_signal(items: list[dict[str, Any]]) -> dict[str, Any]:
    if not items:
        return {}
    title_keywords = (
        "model", "gpt", "gemini", "qwen", "deepseek", "glm", "seed", "kimi", "minimax",
        "hunyuan", "mimo", "ernie", "agent", "coding", "模型", "智能体", "推理", "多模态",
    )

    def score(item: dict[str, Any]) -> tuple[int, float, str]:
        title = str(item.get("title", "")).casefold()
        technical = sum(1 for keyword in title_keywords if keyword in title)
        return technical, float(item.get("confidence") or 0.0), str(item.get("published_at") or "")

    return max(items, key=score)
